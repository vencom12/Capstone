import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_non_technical_doc(filename):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("STITCH-OPT SYSTEM EDGE CASES & WHAT-IF GUIDE")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Non-Technical Business Edition: What is Expected vs. What Actually Happens")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(79, 70, 229) # Indigo
    
    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(16)
    run_intro = intro_p.add_run(
        "This document describes real-world scenarios and 'what-if' situations across the Stitch platform in plain English. "
        "It contrasts what a user naturally expects to happen against what the system actually does behind the scenes to safeguard orders, "
        "finances, and store integrity."
    )
    run_intro.font.name = "Segoe UI"
    run_intro.font.size = Pt(10)
    run_intro.font.color.rgb = RGBColor(51, 65, 85)

    categories = [
        {
            "title": "1. Customer Account & Login Scenarios",
            "desc": "Registration, 1-click Google OAuth, password management, and account deletion.",
            "rows": [
                (
                    "What if a customer tries to register using a Google account that is already registered in Stitch?",
                    "The customer expects not to get an error, and expects to be logged into their existing account smoothly without creating a duplicate mess.",
                    "The system recognizes their Google email, bypasses duplicate registration, automatically signs them into their existing account, and welcomes them back.",
                    "Seamless Sign-In (No crash or duplicate)"
                ),
                (
                    "What if a customer tries to log in with Google, but their account was previously deactivated by management?",
                    "The customer might expect to sneak back in using Google's 1-click button.",
                    "The system intercepts the login immediately and displays: 'Your account has been deactivated. Please contact store management.' Access is blocked.",
                    "Access Denied Safely"
                ),
                (
                    "What if a customer types their email with capital letters (e.g. John@Gmail.com) instead of lowercase?",
                    "The customer expects the system to recognize them regardless of uppercase or lowercase letters.",
                    "The system automatically normalizes all emails to lowercase before checking. Login and password reset work seamlessly.",
                    "Case-Insensitive Matching"
                ),
                (
                    "What if a customer clicks 'Forgot Password' 10 times in a row because their inbox is slow?",
                    "The customer expects the reset email, but shouldn't be allowed to overload the store's email server.",
                    "The system sends the first email and starts a 60-second cooldown timer. Any further clicks show: 'Please wait X seconds before requesting another email.'",
                    "Email Flooding Blocked"
                ),
                (
                    "What if a customer wants to permanently delete their account, but they still have an active order being embroidered?",
                    "The customer expects their private data removed, but the business still needs to fulfill and record the transaction.",
                    "The customer must type their exact username to confirm deletion. The system removes their name, phone, and password, but preserves the anonymous financial order record for auditing.",
                    "Privacy Cleared & Records Intact"
                ),
                (
                    "What if a customer's phone or computer gets stolen while they are logged in?",
                    "The customer expects to change their password on another device and immediately kick the stolen device out.",
                    "Changing the password immediately terminates all active login sessions across all devices, requiring re-login on the next visit.",
                    "Immediate Remote Logout"
                )
            ]
        },
        {
            "title": "2. Shopping, Cart & Checkout Scenarios",
            "desc": "Stock availability, live price changes, discounts, and concurrent buyers.",
            "rows": [
                (
                    "What if two customers try to buy the very last available shirt at the exact same second?",
                    "Both customers expect to get the shirt, but only one can physically be sold.",
                    "The system processes the requests one by one. The first customer gets the item, and the second customer is politely informed: 'Insufficient stock available.'",
                    "Zero Overselling"
                ),
                (
                    "What if a customer leaves an item in their cart for 3 days, and the store increases the price in the meantime?",
                    "The customer expects to see the true current price, and the store cannot afford to sell at an outdated lower price.",
                    "The moment the customer opens the cart or proceeds to checkout, the system re-fetches the live price from the catalog and recalculates the total automatically.",
                    "Live Accurate Pricing"
                ),
                (
                    "What if an item is shown in an A/B price test (Price A vs Price B)?",
                    "The customer expects the price they saw on the screen to be the exact price they are charged at checkout.",
                    "The system locks the assigned test price to that specific customer's session, guaranteeing they pay exactly what was displayed with no sudden surprises.",
                    "Fair Price Consistency"
                ),
                (
                    "What if a customer accidentally clicks 'Place Order' five times because of slow internet?",
                    "The customer expects only one order to be made, not five duplicate charges.",
                    "The order button disables immediately with a loading spinner on the first click, preventing accidental multi-submissions.",
                    "No Duplicate Orders"
                )
            ]
        },
        {
            "title": "3. Payments, GCash & Fake Receipt Scenarios",
            "desc": "Digital wallet payments, screenshot tampering, and payment amounts.",
            "rows": [
                (
                    "What if a customer re-uploads an old GCash payment receipt from last week to pay for a new order?",
                    "A dishonest user expects the system might not notice and approve their order for free.",
                    "The system calculates a digital fingerprint of the image. It spots the duplicate image instantly and rejects it with: 'Duplicate receipt already used.'",
                    "Duplicate Fraud Blocked"
                ),
                (
                    "What if a customer crops or edits an old receipt so the image looks slightly different?",
                    "The user expects to trick the image scanner by changing the pixels.",
                    "The system's AI scanner reads the official GCash Reference Number inside the photo. Because that reference number was already used previously, it flags it as fraud.",
                    "Anti-Replay Fraud Stopped"
                ),
                (
                    "What if a customer uploads a screenshot of a funny meme or a selfie instead of a GCash receipt?",
                    "The user might test if anyone is checking.",
                    "The AI scanner detects that the image is not a valid receipt (confidence score fails) and marks the payment as 'Pending Manual Staff Verification.'",
                    "Fake Uploads Trapped"
                ),
                (
                    "What if a customer's total is ₱1,500, but they only send a receipt for ₱500?",
                    "The customer might hope the order goes through without staff checking the math.",
                    "The system's AI reads the exact amount on the receipt. It sees ₱500 is less than ₱1,500 and marks the order as 'Underpaid / Partial Payment' requiring staff attention.",
                    "Underpayment Flagged"
                ),
                (
                    "What if the internet or AI scanner is temporarily offline when the customer uploads a receipt?",
                    "The customer expects their order not to crash or disappear.",
                    "The system safely saves the receipt image, confirms the upload to the customer, and queues it for manual one-click approval by staff in the Admin Dashboard.",
                    "Zero Order Loss"
                )
            ]
        },
        {
            "title": "4. Custom Embroidery & Workshop Scenarios",
            "desc": "Bring Your Own Garment (BYOG), custom cursive text, and workshop machinery.",
            "rows": [
                (
                    "What if a customer brings their own expensive jacket for embroidery, but refuses to sign the waiver?",
                    "The customer expects to proceed, but the business cannot take liability without legal agreement.",
                    "The checkout system locks the BYOG option until the customer checks and signs the digital damage waiver.",
                    "Liability Protection Enforced"
                ),
                (
                    "What if a customer enters a huge paragraph of text for custom collar embroidery?",
                    "The customer might expect a long essay to fit on a small shirt collar.",
                    "The system enforces a character limit and preview validation, preventing text that exceeds the embroidery hoop boundaries.",
                    "Design Fitting Guaranteed"
                ),
                (
                    "What if an embroidery machine suddenly breaks down in the shop?",
                    "Staff expect to halt orders on that machine without losing track of existing work.",
                    "Staff switch the machine's status to 'Maintenance' in the admin panel. The system immediately shifts pending jobs to other active machines.",
                    "Smooth Machine Reassignment"
                ),
                (
                    "What if a customer pays for a Rush order on a busy weekend?",
                    "The customer expects their order to jump ahead of standard orders.",
                    "The system's AI queuing algorithm assigns high priority score points to Rush orders and automatically moves them to the front of the production line.",
                    "Rush Priority Honored"
                )
            ]
        },
        {
            "title": "5. Admin & Staff Management Scenarios",
            "desc": "Staff permissions, inventory logs, and system controls.",
            "rows": [
                (
                    "What if an admin accidentally tries to deactivate or delete their own admin account?",
                    "An admin might misclick their own row in the user list.",
                    "The system blocks the action immediately: 'You cannot deactivate your own account.'",
                    "Self-Lockout Prevented"
                ),
                (
                    "What if someone tries to delete the only remaining admin in the store?",
                    "A rogue action could leave the business without any administrator.",
                    "The system counts active administrators. If only one remains, deactivation is strictly prohibited.",
                    "Store Management Protected"
                ),
                (
                    "What if someone tries to turn off the Inventory Audit Trail to hide missing thread cones?",
                    "A dishonest employee might try to disable logging before taking inventory.",
                    "The audit trail toggle has been permanently locked to 'ALWAYS ON'. Every single item added, deducted, or modified creates a permanent, non-deletable log entry.",
                    "Tamper-Proof Audit Trail"
                ),
                (
                    "What if an admin exports the entire customer list to Excel / CSV?",
                    "The admin expects customer contact information, but confidential credentials must never be exported.",
                    "The exported CSV only contains business data (name, email, registration date, order count). Passwords and payment credentials are completely excluded.",
                    "Safe Sanitized Export"
                )
            ]
        },
        {
            "title": "6. Connectivity, Device & Browser Scenarios",
            "desc": "Lost Wi-Fi, browser tabs, pop-up blockers, and server restarts.",
            "rows": [
                (
                    "What if a customer's browser blocks pop-ups when clicking 'Continue with Google'?",
                    "The customer expects login to still work even if their browser has a strict pop-up blocker.",
                    "The system detects that the pop-up was blocked and smoothly transitions to full-screen Google login, ensuring they can still sign in.",
                    "Automatic Fallback Login"
                ),
                (
                    "What if an admin logs into the shop on a shared computer, closes the browser window, and walks away?",
                    "Staff sessions must not remain accessible to the next person who opens the computer.",
                    "Staff credentials are configured as temporary browser-session only. Re-opening the browser requires logging in again from scratch.",
                    "Walk-Away Protection Active"
                ),
                (
                    "What if the store's server restarts or has a quick network hiccup while a customer is browsing?",
                    "The customer expects the website not to crash with an ugly error screen.",
                    "The website client automatically retries the request behind the scenes with a short backoff pause, seamlessly absorbing the connection blip.",
                    "Automatic Reconnection"
                ),
                (
                    "What if a customer opens Stitch in two separate browser tabs at the same time?",
                    "The customer expects items added in Tab 1 to appear in Tab 2 without refreshing.",
                    "The system synchronizes shopping cart items across all open tabs in real-time.",
                    "Synchronized Browser Tabs"
                ),
                (
                    "What if an administrator rapidly switches between Customer Reviews, Audit Trail, and Order History tabs?",
                    "The admin expects the screens to appear instantly without staring at 'Loading...' text or wireframe skeleton placeholders on every click.",
                    "The system implements Stale-While-Revalidate (SWR) in-memory caching. It remembers previously fetched records, rendering the entire table in zero milliseconds (<16ms). Behind the scenes, it quietly synchronizes with the server and updates without visual interruption.",
                    "Instant Tab Navigation (SWR)"
                )
            ]
        }
    ]

    for cat in categories:
        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(3)
        run_h2 = h2.add_run(cat["title"])
        run_h2.font.name = "Segoe UI"
        run_h2.font.size = Pt(13)
        run_h2.font.bold = True
        run_h2.font.color.rgb = RGBColor(30, 41, 59)
        
        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_after = Pt(6)
        run_desc = desc_p.add_run(cat["desc"])
        run_desc.font.name = "Segoe UI"
        run_desc.font.size = Pt(9)
        run_desc.font.italic = True
        run_desc.font.color.rgb = RGBColor(100, 116, 139)

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        widths = [Inches(1.8), Inches(2.0), Inches(2.2), Inches(1.2)]
        hdr_cells = table.rows[0].cells
        headers = ["What If... (Scenario)", "What is Expected to Happen", "What Actually Happens in Stitch", "Safety Outcome"]
        for i, text in enumerate(headers):
            hdr_cells[i].width = widths[i]
            set_cell_background(hdr_cells[i], "1E293B")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
            p = hdr_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, (q, exp, act, out) in enumerate(cat["rows"]):
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for i, val in enumerate([q, exp, act, out]):
                row_cells[i].width = widths[i]
                set_cell_background(row_cells[i], bg_color)
                set_cell_margins(row_cells[i], top=90, bottom=90, left=100, right=100)
                p = row_cells[i].paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.name = "Segoe UI"
                run.font.size = Pt(8.5)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                elif i == 3:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(16, 149, 106)
                else:
                    run.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(filename)
    print(f"Generated: {filename}")

def create_technical_doc(filename):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("STITCH-OPT TECHNICAL EDGE CASES & ARCHITECTURE DEFENSES")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Engineering & Security Reference: Protocols, Defenses & Concurrency Handlers")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(14, 165, 233) # Sky Blue
    
    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(16)
    run_intro = intro_p.add_run(
        "This engineering reference documents code-level edge cases, database transaction boundaries, cryptographic defenses, "
        "and concurrency handlers implemented across Express, Next.js, Prisma, PostgreSQL, and Firebase."
    )
    run_intro.font.name = "Segoe UI"
    run_intro.font.size = Pt(10)
    run_intro.font.color.rgb = RGBColor(51, 65, 85)

    categories = [
        {
            "title": "1. Authentication, Sessions & Cryptography",
            "desc": "OAuth protocols, JWT versioning, bcrypt hashing, and database lookups.",
            "rows": [
                (
                    "OAuth Registration Race / Collision",
                    "User registers via Google with an existing account email.",
                    "Controller normalizes email to lower-case and searches via Prisma mode: 'insensitive'. If found, updates verification timestamp and issues signed 30-day JWT cookie without password collision.",
                    "authController.googleAuth"
                ),
                (
                    "Deactivated Account OAuth Bypass",
                    "A soft-disabled user attempts to authenticate via Google OAuth.",
                    "Backend validates `user.isActive === false`. Immediately returns HTTP 403 Forbidden with sanitized error message, terminating token generation.",
                    "authController.googleAuth"
                ),
                (
                    "Cross-Site Request Forgery (CSRF) on Mutations",
                    "Malicious third-party origin submits state-changing POST/PUT request.",
                    "Express csrfToken middleware validates custom `X-CSRF-Token` header and `SameSite=Lax` cookies; invalid tokens abort with HTTP 403.",
                    "server.js & api.ts"
                ),
                (
                    "Token Stolen / Multi-Device Revocation",
                    "User resets credentials or admin forces a global session purge.",
                    "Increments `tokenVersion` in PostgreSQL User record. Future requests fail JWT version check and force immediate logout.",
                    "auth.js & authController.js"
                )
            ]
        },
        {
            "title": "2. Payment Verification & Anti-Replay Cryptographic Defenses",
            "desc": "SHA-256 fingerprinting, OCR reference parsing, and Groq AI vision validation.",
            "rows": [
                (
                    "Image Hash Collision (Duplicate Upload)",
                    "Customer uploads identical GCash screenshot for a second order.",
                    "Server computes crypto SHA-256 hash of image buffer (`imageHash`). Unique constraint on Receipt table rejects duplicate immediately.",
                    "Receipt.imageHash Unique Index"
                ),
                (
                    "Altered / Cropped Image Replay Fraud",
                    "Customer crops receipt border to change image SHA-256 hash.",
                    "Groq AI Vision OCR parses transaction reference number (`referenceId`). Database checks for prior occurrences of `referenceId`.",
                    "Receipt.referenceId Verification"
                ),
                (
                    "Non-Receipt / Corrupted Image Upload",
                    "Customer uploads irrelevant image or corrupted binary.",
                    "Groq AI confidence score drops below `minConfidenceScore` (0.75). Status defaults to `Pending Manual Review` without auto-approval.",
                    "aiReceiptVerifier.js"
                ),
                (
                    "AI Vision API Service Outage",
                    "Groq API returns HTTP 500, 429 rate limit, or socket timeout.",
                    "Try-catch block captures exception, uploads image to storage, and flags receipt for manual staff review in Admin portal.",
                    "aiReceiptVerifier.js Fallback"
                )
            ]
        },
        {
            "title": "3. Concurrency, Inventory & Order Transactions",
            "desc": "Atomic database operations, stock deductions, and race condition prevention.",
            "rows": [
                (
                    "Simultaneous Stock Exhaustion",
                    "Multiple customers checkout the final product inventory unit concurrently.",
                    "Prisma `$transaction` runs atomic decrement. Second transaction triggers condition failure and aborts with rollback.",
                    "checkoutController.js"
                ),
                (
                    "Inventory Audit Log Bypass Attempt",
                    "Admin or automated script attempts to toggle off audit logging.",
                    "InventoryManager permanently enforces `ALWAYS ON` audit logging. Every quantity modification logs immutable `InventoryLog` record.",
                    "utils/inventoryManager.js"
                ),
                (
                    "Price Tampering via Client Payload",
                    "Attacker sends modified `totalAmount` in request JSON body.",
                    "Backend discards client-sent price, queries live database product rates, and computes server-authoritative checkout sum.",
                    "orderController.js"
                )
            ]
        },
        {
            "title": "4. Network, Browser & Framework Resiliency",
            "desc": "Popup blockers, session storage partitioning, and connection timeouts.",
            "rows": [
                (
                    "Browser Popup Blocker Activated",
                    "Browser blocks `signInWithPopup` window.",
                    "Catch block detects `auth/popup-blocked` and gracefully falls back to `signInWithRedirect`, caught by GlobalAuthListener.",
                    "AuthModal.tsx & GlobalAuthListener.tsx"
                ),
                (
                    "Transient Server Disconnection",
                    "Node.js server restarts or brief network blip occurs during fetch.",
                    "Frontend `apiFetch` wrapper executes automatic exponential backoff retry (600ms, then 1200ms) before reporting error.",
                    "frontend/src/lib/api.ts"
                ),
                (
                    "Privileged Tab Abandonment on Shared PC",
                    "Staff leaves browser open and closes device.",
                    "Admin and Employee tokens omit persistent `maxAge` and check `sessionStorage`. Browser exit destroys authentication.",
                    "useAuthStore.ts"
                ),
                (
                    "Dashboard Tab Hopping & Layout Thrashing (SWR Architecture)",
                    "Admin repeatedly navigates between Customer Reviews, Audit Ledger, and Order History tabs.",
                    "Implements module-level Stale-While-Revalidate (SWR) in-memory caching. Returning to visited tabs instantly serves existing dataset (0ms latency, zero skeleton flash) while quietly dispatching an asynchronous background revalidation query. Skeletons and loading spinners are strictly isolated to active user search queries, page resizing, or date filter mutations.",
                    "PanelCustomerReviews.tsx, PanelAuditTrail.tsx, app/admin/page.tsx"
                )
            ]
        }
    ]

    for cat in categories:
        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(14)
        h2.paragraph_format.space_after = Pt(3)
        run_h2 = h2.add_run(cat["title"])
        run_h2.font.name = "Segoe UI"
        run_h2.font.size = Pt(13)
        run_h2.font.bold = True
        run_h2.font.color.rgb = RGBColor(30, 41, 59)
        
        desc_p = doc.add_paragraph()
        desc_p.paragraph_format.space_after = Pt(6)
        run_desc = desc_p.add_run(cat["desc"])
        run_desc.font.name = "Segoe UI"
        run_desc.font.size = Pt(9)
        run_desc.font.italic = True
        run_desc.font.color.rgb = RGBColor(100, 116, 139)

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        widths = [Inches(1.8), Inches(1.8), Inches(2.4), Inches(1.2)]
        hdr_cells = table.rows[0].cells
        headers = ["Technical Edge Case", "Trigger Condition", "Architectural Defense & Resolution", "Code Location"]
        for i, text in enumerate(headers):
            hdr_cells[i].width = widths[i]
            set_cell_background(hdr_cells[i], "0F172A")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=120, right=120)
            p = hdr_cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Segoe UI"
            run.font.size = Pt(9.0)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, (edge, trig, defn, loc) in enumerate(cat["rows"]):
            row_cells = table.add_row().cells
            bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for i, val in enumerate([edge, trig, defn, loc]):
                row_cells[i].width = widths[i]
                set_cell_background(row_cells[i], bg_color)
                set_cell_margins(row_cells[i], top=90, bottom=90, left=100, right=100)
                p = row_cells[i].paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(val)
                run.font.name = "Segoe UI"
                run.font.size = Pt(8.5)
                if i == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(15, 23, 42)
                elif i == 3:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(14, 165, 233)
                else:
                    run.font.color.rgb = RGBColor(51, 65, 85)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.save(filename)
    print(f"Generated: {filename}")

if __name__ == "__main__":
    non_tech = os.path.join(os.getcwd(), "Stitch_System_Edge_Cases_Non_Technical.docx")
    tech = os.path.join(os.getcwd(), "Stitch_System_Edge_Cases_Technical.docx")
    create_non_technical_doc(non_tech)
    create_technical_doc(tech)
